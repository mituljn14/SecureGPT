# pipeline/converters/xlsx_to_md.py
"""
XLSX -> Markdown. One or more real markdown tables per sheet.

Fixes over the old excel_extractor.py:
  1. Merged cells are filled across their whole range before rendering
     (openpyxl only keeps the value in the merge's top-left cell; a
     merged "Q1 FY2020" header spanning 3 columns previously rendered as
     "Q1 FY2020 | | " which loses which columns it covers).
  2. Output is a real markdown table (--- separator row), not bare
     " | "-joined text, and each sheet/block gets a caption line so an
     embedding model has context for what the table even is.
  3. Percentage-formatted cells render as "5.23%" (reading the cell's
     actual number_format) instead of the raw stored fraction
     (0.0522565320665082) -- Excel stores percentages as decimals and
     only the cell format makes them readable.
  4. Plain floats are rounded (4dp) to kill binary-float noise like
     4.39999999999999 -> 4.4, without touching integers.
  5. A merged title/subtitle banner row (common above real data tables --
     every cell in the row/block is the same repeated value from the
     merge-fill) is no longer emitted as its own nonsense 1-row "table".
     It's folded into the caption of the table that follows it instead
     (or kept as a plain line if it's the last thing on the sheet).

Sheets with multiple tables separated by fully-blank rows are still split
into separate tables (handles dashboard-style sheets with several small
tables stacked on one sheet).
"""
import re
from typing import List, Tuple

import openpyxl

from .table_md import rows_to_markdown_table

_PCT_DECIMALS_RE = re.compile(r"0(\.(0+))?%")


def _pct_decimals(number_format: str) -> int:
    m = _PCT_DECIMALS_RE.search(number_format or "")
    if not m or not m.group(2):
        return 0
    return len(m.group(2))


def _format_value(value, number_format: str):
    if isinstance(value, float):
        if number_format and "%" in number_format:
            return f"{value * 100:.{_pct_decimals(number_format)}f}%"
        rounded = round(value, 4)
        # drop a trailing .0 for whole-number floats so 4.0 -> "4", not "4.0"
        return int(rounded) if rounded == int(rounded) else rounded
    return value


def _sheet_grid(ws) -> List[list]:
    """Row/col grid of formatted cell values, with merged-cell ranges
    filled so every cell in a merge carries the top-left value."""
    max_row, max_col = ws.max_row, ws.max_column
    values: List[list] = [[None] * max_col for _ in range(max_row)]
    formats: List[list] = [[None] * max_col for _ in range(max_row)]

    for row in ws.iter_rows():
        for cell in row:
            values[cell.row - 1][cell.column - 1] = cell.value
            formats[cell.row - 1][cell.column - 1] = cell.number_format

    for merged_range in ws.merged_cells.ranges:
        top_value = values[merged_range.min_row - 1][merged_range.min_col - 1]
        top_format = formats[merged_range.min_row - 1][merged_range.min_col - 1]
        for r in range(merged_range.min_row - 1, merged_range.max_row):
            for c in range(merged_range.min_col - 1, merged_range.max_col):
                values[r][c] = top_value
                formats[r][c] = top_format

    return [
        [_format_value(values[r][c], formats[r][c]) for c in range(max_col)]
        for r in range(max_row)
    ]


def _row_banner_value(row: list) -> str | None:
    """A row is a banner (title/subtitle) row, not real tabular data, if
    every non-empty cell in it holds the SAME value -- i.e. it's just a
    merged heading cell that got fill-propagated across the row. Title
    and subtitle rows are typically two DIFFERENT such values, so this
    checks per row, not per block."""
    distinct = {str(v).strip() for v in row if v not in (None, "")}
    if len(distinct) == 1:
        return next(iter(distinct))
    return None


def convert(file_path: str) -> str:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    out = []

    for ws in wb.worksheets:
        grid = _sheet_grid(ws)

        blocks: List[List[list]] = []
        current: List[list] = []
        for row in grid:
            is_blank = all(v is None or str(v).strip() == "" for v in row)
            if is_blank:
                if current:
                    blocks.append(current)
                    current = []
                continue
            current.append(row)
        if current:
            blocks.append(current)

        if not blocks:
            continue

        out.append(f"\n## Sheet: {ws.title}")

        pending_banner = None
        real_table_idx = 0
        for block in blocks:
            # Peel off any leading banner rows (title/subtitle merges) --
            # each can carry a DIFFERENT repeated value, so this walks
            # forward as long as each row individually is a single value.
            banner_parts = []
            idx = 0
            while idx < len(block):
                bv = _row_banner_value(block[idx])
                if bv is None:
                    break
                banner_parts.append(bv)
                idx += 1

            if idx == len(block):
                # every row in this block was a banner row -- no real
                # table here, carry it forward as caption for the next one
                joined = " — ".join(banner_parts)
                pending_banner = (pending_banner + " — " if pending_banner else "") + joined
                continue

            real_table_idx += 1
            header, body = block[idx], block[idx + 1:]
            base_caption = ws.title if len(blocks) == 1 else f"{ws.title} — table {real_table_idx}"
            local_banner = " — ".join(banner_parts) if banner_parts else None
            full_banner = " — ".join(b for b in (pending_banner, local_banner) if b)
            caption = f"{full_banner} ({base_caption})" if full_banner else base_caption
            pending_banner = None

            md_table = rows_to_markdown_table(header, body, caption=caption)
            if md_table:
                out.append("\n" + md_table)

        if pending_banner:  # a banner block with nothing real after it on the sheet
            out.append(f"\n**{pending_banner}**")

    return "\n".join(out).strip()
