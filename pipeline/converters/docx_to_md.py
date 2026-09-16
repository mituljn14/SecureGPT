# pipeline/converters/docx_to_md.py
"""
DOCX -> Markdown, walking paragraphs and tables in document order.

  - "Heading N" styled paragraphs -> "#" * N markdown headings.
  - Tables -> real markdown tables (header + --- separator row) via
    table_md.rows_to_markdown_table, instead of the old " | "-joined text
    with no separator row (which a markdown table detector can't recognize
    and which fed straight into a sentence-splitter that mangled it).
"""
import re
from docx import Document

from .table_md import rows_to_markdown_table

_HEADING_LEVEL_RE = re.compile(r"Heading (\d+)")


def _heading_level(style_name: str) -> int:
    m = _HEADING_LEVEL_RE.match(style_name or "")
    if m:
        return min(int(m.group(1)), 6)
    if style_name == "Title":
        return 1
    return 0


def _table_to_markdown(table) -> str:
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    header, body = rows[0], rows[1:]
    return rows_to_markdown_table(header, body)


def convert(file_path: str) -> str:
    doc = Document(file_path)

    body = doc.element.body
    tables_by_element = {t._tbl: t for t in doc.tables}
    paras_by_element = {p._p: p for p in doc.paragraphs}

    out = []
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]

        if tag == "p":
            para = paras_by_element.get(child)
            if para is None:
                continue
            text = para.text.strip()
            if not text:
                continue
            style = para.style.name if para.style else ""
            level = _heading_level(style)
            if level:
                out.append(f"\n{'#' * level} {text}\n")
            else:
                out.append(text)

        elif tag == "tbl":
            table = tables_by_element.get(child)
            if table is not None:
                md_table = _table_to_markdown(table)
                if md_table:
                    out.append("\n" + md_table + "\n")

    return "\n".join(out).strip()
