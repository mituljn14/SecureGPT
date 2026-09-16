# pipeline/converters/pptx_to_md.py
"""
PPTX -> Markdown, one "## Slide N: <title>" section per slide.

  - Title placeholder -> section heading.
  - Other text frames -> bullet-ish body text, in shape order.
  - Tables -> real markdown tables.
  - Speaker notes -> "### Notes" sub-section (often contains the actual
    substance behind a bare chart/title slide).
"""
from pptx import Presentation

from .table_md import rows_to_markdown_table


def _table_to_markdown(shape) -> str:
    tbl = shape.table
    rows = [[cell.text.strip() for cell in row.cells] for row in tbl.rows]
    if not rows:
        return ""
    header, body = rows[0], rows[1:]
    return rows_to_markdown_table(header, body)


def convert(file_path: str) -> str:
    prs = Presentation(file_path)
    out = []

    for i, slide in enumerate(prs.slides, start=1):
        title = None
        if slide.shapes.title and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()

        out.append(f"\n## Slide {i}" + (f": {title}" if title else ""))

        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_table:
                md_table = _table_to_markdown(shape)
                if md_table:
                    out.append("\n" + md_table)
            elif shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    out.append(text)

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.append(f"\n### Notes\n{notes}")

    return "\n".join(out).strip()
